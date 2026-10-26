# Practice Question #28

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use cell() to access cell B3.
# Print the value of cell B3.
# Use cell() to update the value of cell B3 to 450.
# Save the updated workbook.
# Print a success message.

from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(3, 2) # Use cell() to access cell B3.
print(cell.value) # Print the value of cell B3.

cell.value = 450 # Use cell() to update the value of cell B3 to 450.

wb.save("products_multiple.xlsx") # Save the updated workbook.
print("Updated workbook saved successfully.") # Print a success message.




