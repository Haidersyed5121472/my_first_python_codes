# Practice Question #24

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use cell() to access cell C1.
# Print the value of cell C1.
# Update the value of cell C1 to "Stock".
# Save the updated workbook.
# Print a success message.

from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(row=1, column=3) # Use cell() to access cell C1.
print(cell.value) # Print the value of cell C1.

ws["C1"] = "Stock" # Update the value of cell C1 to "Stock".
wb.save("products_multiple.xlsx") # Save the updated workbook.
print("Workbook updated successfully.") # Print a success message.



