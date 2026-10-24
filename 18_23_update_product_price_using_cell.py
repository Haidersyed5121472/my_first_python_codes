# Practice Question #23

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use cell() to access cell B4.
# Print the value of cell B4.
# Update the value of cell B4 to 175.
# Save the updated workbook.
# Print a success message.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(row=4, column=2) # Use cell() to access cell B4.
print(cell.value) # Print the value of cell B4.

ws["B4"] = 175 # Update the value of cell B4 to 175.

wb.save("products_multiple.xlsx") # Save the updated workbook.
print("Updated workbook saved successfully.") # Print a success message.



