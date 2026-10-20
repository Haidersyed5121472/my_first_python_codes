# Practice Question #15

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use iter_rows() to read cells from row 2 to row 4.
# Read only columns A and B.
# Print the value of each cell.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

for rows in ws.iter_rows(min_row=2, max_row=4, min_col=1, max_col=2): # Use iter_rows() to read rows 2 to 4 and columns A to B.
    for colmn in rows: # Loop through each cell in the current row.
        print(colmn.value) # Print the value of each cell.



