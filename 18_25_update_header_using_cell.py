# Practice Question #25

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use cell() to access cell A1.
# Print the value of cell A1.
# Use cell() to update the value of cell A1 to "Item".
# Save the updated workbook.
# Print a success message.

from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(row=1, column=1) # Use cell() to access cell A1.
print(cell.value) # Print the value of cell A1.

cell.value = "Item" # Use cell() to update the value of cell A1 to "Item".
wb.save("products_multiple.xlsx") # Save the updated workbook.
print("Updated workbook saved successfully.") # Print a success message.


