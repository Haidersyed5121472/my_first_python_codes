# Practice Question #16

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use iter_rows() to read only column B.
# Read all rows containing data.
# Print each cell's value.

from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

for rows in ws.iter_rows(min_col=2, max_col=2): # Use iter_rows() to read only column B.
    for colmn in rows: # Loop through the selected cell in the current row.
        print(colmn.value) # Print each cell's value.



