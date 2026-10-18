# Practice Question #12

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Change the price of "Baby Shoes" from 350 to 400.
# Save the updated workbook.
# Print a success message.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

ws["B3"] = 400 # Change the price of "Baby Shoes" from 350 to 400.

wb.save("products_multiple.xlsx") # Save the updated workbook.
print("Workbook updated successfully.") # Print a success message.


