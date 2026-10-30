
# Practice Question #34

# Import the required library.
# Open the existing "products.xlsx" workbook.
# Select the active worksheet.
# Access cell A1 using cell().
# Align the text in cell A1 to the center horizontally.
# Save the updated workbook.
# Print a success message.


from openpyxl import load_workbook
from openpyxl.styles import Alignment

wb = load_workbook("products.xlsx") # Open the existing "products.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(1, 1) # Access cell A1 using cell().
cell_align = Alignment(horizontal="center") # Align the text in cell A1 to the center horizontally.
cell.alignment = cell_align # Apply the alignment to cell A1.

wb.save("products.xlsx") # Save the updated workbook.
print("Updated workbook saved successfully.") # Print a success message.


