# Practice Question #22

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use cell() to access cell A2.
# Print the value of cell A2.
# Update the value of cell A2 to "Baby Romper".
# Save the updated workbook.
# Print a success message.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(row=2, column=1) # Use cell() to access cell A2.
print(cell.value) # Print the value of cell A2.

ws["A2"] = "Baby Romper" # Update the value of cell A2 to "Baby Romper".
wb.save("products_multiple.xlsx") # Save the updated workbook.

print("Workbook updated successfully.") # Print a success message.




