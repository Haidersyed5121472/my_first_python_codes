
# Practice Question #32

# Import the required library.
# Open the existing "products.xlsx" workbook.
# Select the active worksheet.
# Access cell A1 using cell().
# Change the background color of cell A1 to yellow.
# Save the updated workbook.
# Print a success message.


from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

wb = load_workbook("products.xlsx") # Open the existing "products.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(1, 1) # Access cell A1 using cell().
cell.font = Font(bold=True, size=16)
cell.fill = PatternFill(fgColor="FFFF00", fill_type="solid") # Change the background color of cell A1 to yellow.

wb.save("products.xlsx") # Save the updated workbook.
print("Updated workbook saved successfully.") # Print a success message.




