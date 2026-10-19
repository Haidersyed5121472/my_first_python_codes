# Practice Question #14

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use iter_rows() with values_only=True.
# Loop through all rows.
# Print each row's values.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

for rows in ws.iter_rows(values_only=True): # Use iter_rows() with values_only=True.
    for cell in rows: # Loop through the values in the current row.
        print(cell) # Print each row's values.



