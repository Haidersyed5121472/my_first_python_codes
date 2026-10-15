# Practice Question #6

# Import the required library.
# Open the existing "products.xlsx" workbook.
# Select the active worksheet.
# Print the total number of rows.
# Print the total number of columns.


from openpyxl import load_workbook

workbook = load_workbook("products.xlsx") # Open the existing "products.xlsx" workbook.
sheet = workbook.active # Select the active worksheet.

print(sheet.max_row) # Print the total number of rows.
print(sheet.max_column) # Print the total number of columns.



