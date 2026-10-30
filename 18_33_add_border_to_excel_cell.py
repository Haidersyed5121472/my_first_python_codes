
# Practice Question #33

# Import the required library.
# Open the existing "products.xlsx" workbook.
# Select the active worksheet.
# Access cell A1 using cell().
# Create a thin black border for all four sides of the cell.
# Apply the border to cell A1.
# Save the updated workbook.
# Print a success message.

from openpyxl import load_workbook
from openpyxl.styles import Border, Side

wb = load_workbook("products.xlsx") # Open the existing "products.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(1, 1) # Access cell A1 using cell().
thin_black = Side(style="thin", color="000000") # Create a thin black border side.
cell.border = Border(
    left=thin_black,
    right=thin_black,
    top=thin_black,
    bottom=thin_black
) # Apply the border to cell A1.

wb.save("products.xlsx") # Save the updated workbook.
print("Updated workbook saved successfully.") # Print a success message.



