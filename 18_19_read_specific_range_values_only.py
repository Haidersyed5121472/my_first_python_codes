# Practice Question #19

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use iter_rows() to read rows 2 to 4 and columns A to B.
# Use values_only=True to get the direct values of each cell.
# Use nested for loops to access each value.
# Print each value.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

# Use iter_rows() to read rows 2 to 4 and columns A to B.
# Use values_only=True to get the direct values of each cell.
for rows in ws.iter_rows(min_row=2, max_row=4, min_col=1, max_col=2, values_only=True):
    for column in rows: # Use nested for loops to access each value.
        print(column) # Print each value.



