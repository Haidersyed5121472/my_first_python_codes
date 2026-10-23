# Practice Question #21

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use cell() to access cell B2.
# Print the value of cell B2.
# Update the value of cell B2 to 650.
# Save the updated workbook.
# Print a success message.

from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(row=2, column=2) # Use cell() to access cell B2.
print(cell.value) # Print the value of cell B2.

ws["B2"] = 650 # Update the value of cell B2 to 650.
wb.save("products_multiple.xlsx") # Save the updated workbook.

print("Workbook updated successfully.") # Print a success message.

