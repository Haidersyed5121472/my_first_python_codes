# Practice Question #17

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use max_row and max_column to get the worksheet's dimensions.
# Use iter_rows() to read all rows and columns containing data.
# Print the value of each cell.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

for rows in ws.iter_rows(max_row=ws.max_row, max_col=ws.max_column): # Use iter_rows() to read all rows and columns containing data.
    for colmn in rows: # Loop through each cell in the current row.
        print(colmn.value) # Print the value of each cell.



