# Practice Question #27

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use cell() to access cell A3.
# Print the value of cell A3.
# Use cell() to update the value of cell A3 to "Baby Shoes Premium".
# Save the updated workbook.
# Print a success message.

from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

cell = ws.cell(3, 1) # Use cell() to access cell A3.
print(cell.value) # Print the value of cell A3.

cell.value = "Baby Shoes Premium" # Use cell() to update the value of cell A3 to "Baby Shoes Premium".

wb.save("products_multiple.xlsx") # Save the updated workbook.
print("Updated workbook saved successfully.") # Print a success message.





