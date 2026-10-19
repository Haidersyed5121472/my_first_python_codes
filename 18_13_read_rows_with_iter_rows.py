# Practice Question #13

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use iter_rows() to read all rows.
# Print the value of each cell.

from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

for rows in ws.iter_rows(): # Use iter_rows() to read all rows.
    for cell in rows: # Use a for loop to go through the cells in the current row.
        print(cell.value) # Print the value of each cell.



