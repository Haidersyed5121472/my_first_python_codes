# Practice Question #18

# Import the required library.
# Open the existing "products_multiple.xlsx" workbook.
# Select the active worksheet.
# Use iter_rows() with values_only=True to read all rows and columns containing data.
# Use nested for loops to access each value.
# Print the value of each cell.


from openpyxl import load_workbook

wb = load_workbook("products_multiple.xlsx") # Open the existing "products_multiple.xlsx" workbook.
ws = wb.active # Select the active worksheet.

for rows in ws.iter_rows(values_only=True): # Use iter_rows() with values_only=True to read all rows and columns containing data.
    for column in rows: # Use nested for loops to access each value.
        print(column) # Print the value of each cell.



